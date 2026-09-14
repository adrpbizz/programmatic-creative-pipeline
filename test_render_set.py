"""Run with: python -m unittest -v"""

import json
import tempfile
import unittest
from pathlib import Path

from PIL import Image
from pptx import Presentation

from render_set import F, autofit, export_pptx, main, read_rows, render_png, resolve, tw


class PipelineTests(unittest.TestCase):
    def test_measure_fit_and_native_exports(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            Image.new("RGB", (80, 40), "#00FF00").save(root / "photo.png")
            template = {
                "version": 1, "width": 400, "height": 400, "background": "white",
                "elements": [
                    dict(kind="rect", x=10, y=10, width=30, height=20, color="#FF0000"),
                    dict(kind="image", x=300, y=300, width=50, height=60, src="photo.png"),
                    dict(kind="text", x=40, y=80, width=320, height=50,
                         text="$title", size=80, tracking=3, font="auto-bold"),
                ],
            }
            spec = resolve(template, {"title": "AV jé!"}, root)
            text = spec["elements"][-1]
            font = F("auto-bold", text["size"])
            self.assertLessEqual(tw(text["text"], font, 3), 320)
            self.assertLess(text["size"], 80)
            self.assertAlmostEqual(text["ink_box"][0] + text["ink_box"][2] / 2, 200)
            self.assertEqual(tw("A", font, 30), tw("A", font))
            self.assertEqual(autofit("", "auto", 8, 50, min_size=8).size, 8)
            with self.assertRaises(ValueError):
                autofit("cannot fit", "auto", 10, 1)

            render_png(spec, root / "creative.png", root)
            with Image.open(root / "creative.png") as im:
                self.assertEqual(im.size, (400, 400))
                self.assertEqual(im.getpixel((10, 10)), (255, 0, 0))
                self.assertEqual(im.getpixel((40, 10)), (255, 255, 255))
                self.assertEqual(im.getpixel((325, 325)), (0, 255, 0))
                # Actual ink must stay in its safe box and center on the axis.
                ink = im.crop((0, 60, 400, 150)).convert("L").point(lambda p: 255 - p).getbbox()
                self.assertIsNotNone(ink)
                self.assertGreaterEqual(ink[0], 40)
                self.assertLessEqual(ink[2], 360)
                self.assertAlmostEqual((ink[0] + ink[2]) / 2, 200, delta=2)
                self.assertGreaterEqual(ink[1] + 60, 80)
                self.assertLessEqual(ink[3] + 60, 130)

            export_pptx([spec, spec], root / "creative.pptx", root)
            prs = Presentation(root / "creative.pptx")
            self.assertEqual(len(prs.slides), 2)
            self.assertEqual(len(prs.slides[0].shapes), 3)
            shape = prs.slides[0].shapes[-1]
            self.assertEqual(shape.text, "AV jé!")
            run = shape.text_frame.paragraphs[0].runs[0]
            self.assertTrue(run.font.bold)
            self.assertEqual(run._r.get_or_add_rPr().get("spc"), "540")
            self.assertNotIn(str(root), json.dumps(spec))

    def test_batch_errors_leave_no_output_and_never_overwrite(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            template = root / "template.json"
            template.write_text(json.dumps({
                "version": 1, "width": 200, "height": 200,
                "elements": [dict(kind="text", x=10, y=10, width=180, height=50,
                                  text="$title", size=30, min_size=20)],
            }), encoding="utf-8")
            csv_path = root / "rows.csv"
            output = root / "output"
            csv_path.write_text("id,title\ngood,Hello\nbad," + "W" * 200, encoding="utf-8")
            with self.assertRaises(SystemExit):
                main([str(template), "--data", str(csv_path), "--out", str(output)])
            self.assertFalse(output.exists())
            csv_path.write_text("id,title\ngood,Hello\nsecond,Bonjour", encoding="utf-8")
            main([str(template), "--data", str(csv_path), "--out", str(output), "--pptx"])
            self.assertEqual(len(list(output.iterdir())), 5)
            before = (output / "good.png").read_bytes()
            with self.assertRaises(SystemExit):
                main([str(template), "--data", str(csv_path), "--out", str(output)])
            self.assertEqual((output / "good.png").read_bytes(), before)
            for content in ("id,title\n../escape,Hello", "id,title\nA,Hello\na,World",
                            "id,title\nCON,Hello", "id,title\nx", "id,id\nx,x"):
                csv_path.write_text(content, encoding="utf-8")
                with self.assertRaises(ValueError):
                    read_rows(csv_path)

    def test_validation(self):
        template = {"version": 1, "width": 100, "height": 100, "elements": []}
        with self.assertRaises(ValueError):
            resolve({**template, "width": float("nan")})
        for element in (
            dict(kind="text", x=0, y=0, width=100, height=100, text="$missing"),
            dict(kind="text", x=0, y=0, width=100, height=100, text="a\nb"),
            dict(kind="rect", x=90, y=0, width=20, height=20),
            dict(kind="rect", x=0, y=0, width=20, height=20, colour="red"),
            dict(kind="rect", x=0, y=0, width=20, height=20, color=123),
            dict(kind="image", x=0, y=0, width=20, height=20, src="../outside.png"),
        ):
            with self.assertRaises(ValueError):
                resolve({**template, "elements": [element]})


if __name__ == "__main__":
    unittest.main()
