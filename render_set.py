"""Render local JSON templates to PNG, resolved JSON, and optional editable PPTX."""

import argparse
import csv
import json
import math
import os
import re
import tempfile
from functools import lru_cache
from pathlib import Path
from string import Template

from PIL import Image, ImageColor, ImageDraw, ImageFont, ImageOps


def number(value, label, minimum=0):
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        raise ValueError(f"{label} must be a number")
    if not math.isfinite(value) or value < minimum:
        raise ValueError(f"{label} must be finite and >= {minimum}")
    return value


def color(value):
    """Normalize an opaque Pillow color to portable #RRGGBB."""
    if not isinstance(value, str):
        raise ValueError("Color must be a string, such as #123456")
    rgb = ImageColor.getrgb(value)
    if len(rgb) != 3:
        raise ValueError("Use opaque RGB colors, such as #123456")
    return "#" + "".join(f"{c:02X}" for c in rgb)


@lru_cache(maxsize=256)
def F(fn, size):
    """Load a TrueType/OpenType font. Explicit paths never silently fall back."""
    windows_fonts = Path(os.environ.get("WINDIR", "C:/Windows")) / "Fonts"
    names = {
        "auto": ("DejaVuSans.ttf", str(windows_fonts / "arial.ttf"), "/System/Library/Fonts/Supplemental/Arial.ttf"),
        "auto-bold": ("DejaVuSans-Bold.ttf", str(windows_fonts / "arialbd.ttf"), "/System/Library/Fonts/Supplemental/Arial Bold.ttf"),
    }.get(str(fn), (str(fn),))
    for name in names:
        try:
            return ImageFont.truetype(name, size)
        except OSError:
            pass
    raise ValueError(f"Font not found: {fn}. Supply a local .ttf/.otf path in the template.")


def ink_bounds(t, f):
    """Measure visible pixels; font advance boxes can include side bearings."""
    mask, (x, y) = f.getmask2(t, anchor="la")
    b = mask.getbbox()
    return (x + b[0], y + b[1], x + b[2], y + b[3]) if b else None


def text_bounds(t, f, tr=0):
    """Measure the same glyph positions we draw, including bearings."""
    if not t:
        return (0, 0, 0, 0)
    if tr == 0:
        return ink_bounds(t, f) or (0, 0, 0, 0)
    # ponytail: tracking is for simple scripts; use zero for joined/complex text.
    pen = 0
    boxes = []
    for c in t:
        b = ink_bounds(c, f)
        if b:
            left, top, right, bottom = b
            boxes.append((pen + left, top, pen + right, bottom))
        pen += f.getlength(c) + tr
    if not boxes:
        return (0, 0, 0, 0)
    return (min(b[0] for b in boxes), min(b[1] for b in boxes),
            max(b[2] for b in boxes), max(b[3] for b in boxes))


def tw(t, f, tr=0):
    b = text_bounds(t, f, tr)
    return b[2] - b[0]


def autofit(t, fn, start, mw, tr=0, *, min_size=10, max_height=None):
    """Return a font that fits; raise when even min_size would overflow."""
    number(start, "size", 1)
    number(min_size, "min_size", 1)
    number(mw, "width", 1)
    number(tr, "tracking")
    if max_height is not None:
        number(max_height, "height", 1)
    if int(start) != start or int(min_size) != min_size or min_size > start or start > 8192:
        raise ValueError("Font sizes must be integers with 1 <= min_size <= size <= 8192")
    for z in range(int(start), int(min_size) - 1, -1):
        f = F(str(fn), z)
        b = text_bounds(t, f, tr)
        if b[2] - b[0] <= mw and (max_height is None or b[3] - b[1] <= max_height):
            return f
    raise ValueError(f"Text cannot fit at minimum size {min_size}: {t[:80]!r}")


def asset_path(reference, base_dir):
    path = Path(reference)
    if path.is_absolute():
        raise ValueError("Use asset paths relative to the template directory")
    root = Path(base_dir).resolve()
    path = (root / path).resolve()
    if not path.is_relative_to(root):
        raise ValueError("Asset paths must stay inside the template directory")
    if not path.is_file():
        raise ValueError(f"Asset not found: {reference}")
    return path


def font_reference(reference, base_dir):
    return reference if reference in ("auto", "auto-bold") else str(asset_path(reference, base_dir))


def resolve(template, data=None, base_dir=Path(".")):
    """Resolve substitutions and text sizes once for both export backends."""
    if not isinstance(template, dict):
        raise ValueError("Template must be a JSON object")
    if template.get("version") != 1:
        raise ValueError("Template version must be 1")
    width = number(template["width"], "canvas width", 1)
    height = number(template["height"], "canvas height", 1)
    if int(width) != width or int(height) != height or max(width, height) > 8192:
        raise ValueError("Canvas dimensions must be integers between 1 and 8192")
    if width * height > 32_000_000:
        raise ValueError("Canvas exceeds the 32 megapixel limit")
    spec = dict(version=1, width=int(width), height=int(height),
                background=color(template.get("background", "white")), elements=[])
    if not isinstance(template["elements"], list):
        raise ValueError("elements must be an array")
    for index, source in enumerate(template["elements"]):
        try:
            kind = source["kind"]
            if kind not in ("text", "rect", "ellipse", "image"):
                raise ValueError(f"Unknown element kind: {kind}")
            common = {"kind", "x", "y", "width", "height"}
            extras = {"text": {"text", "font", "size", "min_size", "tracking", "color", "align"},
                      "rect": {"color"}, "ellipse": {"color"}, "image": {"src"}}[kind]
            unknown = source.keys() - common - extras
            if unknown:
                raise ValueError(f"Unknown fields: {', '.join(sorted(unknown))}")
            e = dict(kind=kind)
            for key in ("x", "y", "width", "height"):
                e[key] = number(source[key], key, 1 if key in ("width", "height") else 0)
            if e["x"] + e["width"] > width or e["y"] + e["height"] > height:
                raise ValueError("Element extends beyond the canvas")
            if kind == "image":
                e["src"] = Template(source["src"]).substitute(data or {})
                path = asset_path(e["src"], base_dir)
                with Image.open(path) as im:
                    im.verify()
            else:
                e["color"] = color(source.get("color", "black"))
            if kind == "text":
                e["text"] = Template(source["text"]).substitute(data or {})
                if any(ord(c) < 32 for c in e["text"]):
                    raise ValueError("Use one text element per line; control characters are unsupported")
                e["tracking"] = number(source.get("tracking", 0), "tracking")
                e["font"] = source.get("font", "auto")
                e["align"] = source.get("align", "center")
                if e["align"] not in ("left", "center", "right"):
                    raise ValueError("align must be left, center, or right")
                f = autofit(e["text"], font_reference(e["font"], base_dir),
                            source.get("size", 64), e["width"], e["tracking"],
                            min_size=source.get("min_size", 10), max_height=e["height"])
                b = text_bounds(e["text"], f, e["tracking"])
                ink_width = b[2] - b[0]
                offset = {"left": 0, "center": (e["width"] - ink_width) / 2,
                          "right": e["width"] - ink_width}[e["align"]]
                family, style = f.getname()
                e.update(size=f.size, font_family=family, bold="bold" in style.lower(),
                         italic=any(s in style.lower() for s in ("italic", "oblique")),
                         cx=e["x"] + e["width"] / 2,
                         ink_box=[e["x"] + offset, e["y"], ink_width, b[3] - b[1]])
            spec["elements"].append(e)
        except (KeyError, TypeError, ValueError, OSError) as exc:
            raise ValueError(f"Element {index + 1}: {exc}") from exc
    return spec


def fitted_image(e, base_dir):
    with Image.open(asset_path(e["src"], base_dir)) as im:
        return ImageOps.fit(ImageOps.exif_transpose(im).convert("RGBA"),
                            (max(1, round(e["width"])), max(1, round(e["height"]))),
                            method=Image.Resampling.LANCZOS)


def render_png(spec, path, base_dir=Path(".")):
    image = Image.new("RGB", (spec["width"], spec["height"]), spec["background"])
    d = ImageDraw.Draw(image)
    for e in spec["elements"]:
        x, y, w, h = (e[k] for k in ("x", "y", "width", "height"))
        if e["kind"] == "image":
            asset = fitted_image(e, base_dir)
            image.paste(asset, (round(x), round(y)), asset)
        elif e["kind"] in ("rect", "ellipse"):
            draw = d.rectangle if e["kind"] == "rect" else d.ellipse
            draw((x, y, x + w - 1, y + h - 1), fill=e["color"])
        else:
            f = F(font_reference(e["font"], base_dir), e["size"])
            b = text_bounds(e["text"], f, e["tracking"])
            pen, top = e["ink_box"][0] - b[0], y - b[1]
            if e["tracking"] == 0:
                d.text((pen, top), e["text"], font=f, fill=e["color"], anchor="la")
            else:
                for c in e["text"]:
                    d.text((pen, top), c, font=f, fill=e["color"], anchor="la")
                    pen += f.getlength(c) + e["tracking"]
    image.save(path, "PNG")


def export_pptx(specs, path, base_dir=Path(".")):
    """Native shapes and text, with images as individually editable pictures."""
    from io import BytesIO
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_AUTO_SIZE, MSO_ANCHOR, PP_ALIGN
    from pptx.util import Inches, Pt

    prs = Presentation()
    if not specs:
        raise ValueError("At least one creative is required")
    # One shared scale preserves pixels-to-points and stays inside PPTX size limits.
    scale = 10 / max(specs[0]["width"], specs[0]["height"])
    emu = lambda px: Inches(px * scale)
    prs.slide_width, prs.slide_height = emu(specs[0]["width"]), emu(specs[0]["height"])
    for spec in specs:
        if (spec["width"], spec["height"]) != (specs[0]["width"], specs[0]["height"]):
            raise ValueError("All slides in one PPTX must have the same dimensions")
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor.from_string(spec["background"][1:])
        for e in spec["elements"]:
            x, y, w, h = (emu(e[k]) for k in ("x", "y", "width", "height"))
            if e["kind"] == "image":
                stream = BytesIO()
                fitted_image(e, base_dir).save(stream, "PNG")
                stream.seek(0)
                slide.shapes.add_picture(stream, x, y, w, h)
            elif e["kind"] in ("rect", "ellipse"):
                kind = MSO_SHAPE.RECTANGLE if e["kind"] == "rect" else MSO_SHAPE.OVAL
                shape = slide.shapes.add_shape(kind, x, y, w, h)
                shape.fill.solid()
                shape.fill.fore_color.rgb = RGBColor.from_string(e["color"][1:])
                shape.line.fill.background()
            else:
                shape = slide.shapes.add_textbox(x, y, w, max(h, emu(e["size"] * 1.5)))
                tf = shape.text_frame
                tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
                tf.word_wrap = False
                tf.auto_size = MSO_AUTO_SIZE.NONE
                tf.vertical_anchor = MSO_ANCHOR.TOP
                p = tf.paragraphs[0]
                p.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER,
                               "right": PP_ALIGN.RIGHT}[e["align"]]
                p.space_before = p.space_after = Pt(0)
                run = p.add_run()
                run.text = e["text"]
                run.font.name = e["font_family"]
                run.font.size = Pt(e["size"] * scale * 72)
                run.font.bold, run.font.italic = e["bold"], e["italic"]
                run.font.color.rgb = RGBColor.from_string(e["color"][1:])
                # python-pptx has no public tracking setter; OOXML uses 1/100 pt.
                run._r.get_or_add_rPr().set("spc", str(round(e["tracking"] * scale * 7200)))
        slide.notes_slide.notes_text_frame.text = "Generated by Programmatic Creative Pipeline. Fonts are not embedded."
    prs.save(path)


def read_rows(path):
    if path is None:
        return [{"id": "creative"}]
    with path.open(encoding="utf-8-sig", newline="") as stream:
        reader = csv.DictReader(stream)
        if not reader.fieldnames or "id" not in reader.fieldnames:
            raise ValueError("CSV needs an id column")
        if len(reader.fieldnames) != len(set(reader.fieldnames)):
            raise ValueError("CSV has duplicate column names")
        rows = list(reader)
    if not rows:
        raise ValueError("CSV has no data rows")
    seen = set()
    for row in rows:
        if None in row or None in row.values():
            raise ValueError("CSV row length does not match the header")
        name = row["id"]
        if not re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9_-]{0,79}", name):
            raise ValueError(f"Invalid id {name!r}: use letters, numbers, underscores or hyphens")
        if name.upper() in {"CON", "PRN", "AUX", "NUL", *(f"COM{i}" for i in range(1, 10)), *(f"LPT{i}" for i in range(1, 10))}:
            raise ValueError(f"Reserved filename: {name}")
        if name.casefold() in seen:
            raise ValueError(f"Duplicate id: {name}")
        seen.add(name.casefold())
    return rows


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("template", type=Path, help="JSON template")
    parser.add_argument("--data", type=Path, help="UTF-8 CSV with id and substitution columns")
    parser.add_argument("--out", type=Path, default=Path("output"), help="New output directory")
    parser.add_argument("--pptx", action="store_true", help="Also create editable creatives.pptx")
    args = parser.parse_args(argv)
    try:
        if args.out.exists():
            raise ValueError(f"Output already exists: {args.out}. Choose a new --out directory.")
        template = json.loads(args.template.read_text(encoding="utf-8-sig"))
        rows = read_rows(args.data)
        base = args.template.resolve().parent
        specs = [resolve(template, row, base) for row in rows]
        args.out.parent.mkdir(parents=True, exist_ok=True)
        # Stage the entire batch: a bad row/export never leaves a half-finished set.
        with tempfile.TemporaryDirectory(prefix=".pcp-", dir=args.out.parent) as temporary:
            stage = Path(temporary) / "set"
            stage.mkdir()
            for row, spec in zip(rows, specs):
                render_png(spec, stage / f"{row['id']}.png", base)
                (stage / f"{row['id']}.json").write_text(json.dumps(spec, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
            if args.pptx:
                export_pptx(specs, stage / "creatives.pptx", base)
            if args.out.exists():
                raise ValueError("Output appeared during rendering; choose a new directory")
            stage.rename(args.out)
        print(f"Rendered {len(specs)} creative(s) to {args.out}")
    except (OSError, ValueError, KeyError, TypeError, ImportError, csv.Error) as exc:
        parser.exit(1, f"Error: {exc}\n")


if __name__ == "__main__":
    main()
